from docx import Document
d = Document()
d.add_heading('Appendix 2 - Service Terms', 1)
p = d.add_paragraph('The Supplier shall deliver the ')
r = p.add_run('Services'); r.bold = True
p.add_run(' described in Schedule A within ')
r2 = p.add_run('thirty (30) days'); r2.italic = True
p.add_run(' of the Effective Date.')
p2 = d.add_paragraph('Payment terms are net sixty (60) days from receipt of a valid invoice.')
d.add_comment(p2.runs, text='Customer proposes net 60; we need net 30.', author='Alex Customer', initials='AC')
d.add_heading('Liability', 2)
d.add_paragraph('Neither party shall be liable for indirect or consequential damages.')
t = d.add_table(rows=2, cols=2); t.style = 'Table Grid'
t.cell(0,0).text = 'Fee'; t.cell(0,1).text = 'Amount'
t.cell(1,0).text = 'Setup'; t.cell(1,1).text = '$10,000'
d.add_paragraph('This Appendix forms part of the Agreement.')
d.save('customer.docx')
